"""Тесты на найденные и исправленные ошибки.

Каждый тест здесь воспроизводит конкретный баг, который уже был в коде.
Если однажды они снова начнут падать — значит, ошибка вернулась.

Запуск: python tests/regressions.py
"""

from __future__ import annotations

import asyncio
import io
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from aiogram import Bot  # noqa: E402
from aiogram.client.default import DefaultBotProperties  # noqa: E402
from aiogram.enums import ParseMode  # noqa: E402
from aiogram.fsm.storage.base import StorageKey  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from flow import BOT_ID, CHAT_ID, USER_ID, FakeSession, cb, contact_msg, msg  # noqa: E402

from app import db, export, texts  # noqa: E402
from app.catalog import CHAIR_NORD  # noqa: E402
from app.config import Settings  # noqa: E402
from app.keyboards import cancel_kb, skip_kb  # noqa: E402
from app.utils import SHOP_TZ, now  # noqa: E402
from bot import create_dispatcher  # noqa: E402

passed = 0
failed: list[str] = []


def check(name: str, condition: bool, detail: object = "") -> None:
    global passed
    if condition:
        passed += 1
        print(f"  ok   {name}")
    else:
        failed.append(f"{name} — {detail}")
        print(f"  FAIL {name} — {detail}")


def build(tmp: Path):
    """Диспетчер создаётся ОДИН раз на весь запуск.

    Роутеры — модульные объекты, и aiogram запрещает подключать один и тот же
    роутер ко второму диспетчеру. Поэтому сценарии изолируются не новым
    диспетчером, а чистой базой и сброшенным состоянием FSM.
    """
    settings = Settings(
        token="0:TEST",
        admin_ids=(USER_ID,),
        viewer_url="https://example.com/3d",
        manager_url="https://t.me/example",
        shop_name="Мебельная мастерская NORD",
        db_path=tmp / "reg.sqlite3",
        assets_dir=tmp / "assets",
    )
    session = FakeSession()
    bot = Bot(
        token="0:TEST",
        session=session,
        default=DefaultBotProperties(parse_mode=ParseMode.HTML),
    )
    return session, bot, create_dispatcher(settings, throttle_rate=0.0)


async def reset(dp, session: FakeSession, tmp: Path, name: str) -> None:
    """Чистая база + пустое состояние: сценарии не влияют друг на друга."""
    await db.init_db(tmp / f"{name}.sqlite3")
    key = StorageKey(bot_id=BOT_ID, chat_id=CHAT_ID, user_id=USER_ID)
    await dp.storage.set_state(key, None)
    await dp.storage.set_data(key, {})
    session.reset()


ORDER_FORM = {
    "name": "Иван Иванов",
    "phone": "+79001112233",
    "city": "Москва",
    "address": "ул. Павла Корчагина, д. 22А, кв. 5",
    "comment": "",
}


async def main() -> None:
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        await db.init_db(tmp / "reg.sqlite3")
        session, bot, dp = build(tmp)

        async def feed(update) -> None:
            await dp.feed_update(bot, update)

        async def reach_confirm() -> None:
            await feed(cb("cart:checkout:0"))
            await feed(msg("Бюньямин Каткат"))
            await feed(contact_msg("89001112233"))
            await feed(msg("Москва"))
            await feed(msg("ул. Павла Корчагина, д. 22А, кв. 5"))
            await feed(msg(texts.BTN_SKIP))

        # ------------------------------------------------------------------ 1
        print("\n[1] Корзину нельзя менять во время оформления (была критическая ошибка)")
        await reset(dp, session, tmp, "s1")

        await feed(msg("/start"))
        await feed(msg(texts.BTN_CONFIGURE))
        await feed(cb("cfg:summary:chair-nord:-:-"))
        await feed(cb("cfg:add:chair-nord:-:-"))
        check("в корзине 1 позиция", len(await db.get_cart(USER_ID)) == 1)

        await reach_confirm()
        confirm_screen = session.texts()[-1]
        check("экран подтверждения показан", "Проверьте заказ" in confirm_screen)
        check("на экране 42 900", "42" in confirm_screen, confirm_screen[-160:])

        # Старая кнопка «В корзину» из карточки выше: раньше она срабатывала
        # прямо в состоянии подтверждения и меняла сумму уже согласованного заказа.
        cart_id = (await db.get_cart(USER_ID))[0].id
        session.reset()
        await feed(cb("cfg:add:chair-nord:-:-"))
        total_qty = sum(item.qty for item in await db.get_cart(USER_ID))
        check("кнопка конфигуратора заблокирована в оформлении", total_qty == 1, total_qty)

        session.reset()
        await feed(cb(f"cart:inc:{cart_id}"))
        total_qty = sum(item.qty for item in await db.get_cart(USER_ID))
        check("кнопка «+» корзины заблокирована в оформлении", total_qty == 1, total_qty)

        session.reset()
        await feed(cb("chk:confirm"))
        orders = await db.user_orders(USER_ID)
        check("заказ создан", len(orders) == 1, len(orders))
        check(
            "сумма заказа == сумме на экране (42 900)",
            bool(orders) and orders[0].total == 42_900,
            orders[0].total if orders else None,
        )

        # ------------------------------------------------------------------ 2
        print("\n[2] Изменение корзины в обход кнопок (второе устройство)")
        await reset(dp, session, tmp, "s2")

        await feed(msg("/start"))
        await feed(cb("cfg:add:chair-nord:-:-"))
        await reach_confirm()

        # Имитируем второе устройство: пишем в базу мимо хендлеров.
        other = CHAIR_NORD.normalize({"fabric": "indigo"})
        await db.add_to_cart(USER_ID, CHAIR_NORD.id, other, CHAIR_NORD.price(other))

        session.reset()
        await feed(cb("chk:confirm"))
        check("заказ НЕ создан по неподтверждённой сумме", not await db.user_orders(USER_ID))
        check(
            "клиент предупреждён об изменении",
            any("изменилась" in text for text in session.texts()),
            session.texts(),
        )
        check(
            "показан новый экран подтверждения",
            any("Проверьте заказ" in text for text in session.texts()),
        )

        session.reset()
        await feed(cb("chk:confirm"))
        orders = await db.user_orders(USER_ID)
        check("после нового согласия заказ создан", len(orders) == 1, len(orders))
        check(
            "и уже на новую сумму",
            bool(orders) and orders[0].total == 42_900 + CHAIR_NORD.price(other),
            orders[0].total if orders else None,
        )

        # ------------------------------------------------------------------ 3
        print("\n[3] Команда внутри оформления не становится именем клиента")
        await reset(dp, session, tmp, "s3")

        await feed(msg("/start"))
        await feed(cb("cfg:add:chair-nord:-:-"))
        await feed(cb("cart:checkout:0"))

        session.reset()
        await feed(msg("/start"))
        check(
            "/start выводит из сценария",
            any("отменено" in text for text in session.texts()),
            session.texts(),
        )

        await feed(cb("cart:checkout:0"))
        session.reset()
        await feed(msg("/help"))
        check(
            "/help не принят как имя — вопрос повторён",
            any("Как к вам обращаться" in text for text in session.texts()),
            session.texts(),
        )

        # ------------------------------------------------------------------ 4
        print("\n[4] Корзина не может вырасти до неоткрываемого размера")
        await reset(dp, session, tmp, "s4")

        added = 0
        for fabric in (option.id for option in CHAIR_NORD.group("fabric").options):
            for frame in (option.id for option in CHAIR_NORD.group("frame").options):
                config = CHAIR_NORD.normalize({"fabric": fabric, "frame": frame})
                if await db.add_to_cart(
                    USER_ID, CHAIR_NORD.id, config, CHAIR_NORD.price(config)
                ):
                    added += 1

        items = await db.get_cart(USER_ID)
        check(
            f"добавлено ровно {db.MAX_POSITIONS} позиций из 24 попыток",
            len(items) == db.MAX_POSITIONS,
            len(items),
        )
        check("лишние отклонены", added == db.MAX_POSITIONS, added)

        cart_text = texts.cart(
            items,
            {CHAIR_NORD.id: CHAIR_NORD.name},
            {item.id: CHAIR_NORD.config_line(item.config) for item in items},
        )
        check(
            f"полная корзина влезает в сообщение ({len(cart_text)} знаков)",
            len(cart_text) <= texts.MESSAGE_LIMIT,
            len(cart_text),
        )

        # ------------------------------------------------------------------ 5
        print("\n[5] «Мои заказы» не ломаются на длинной истории")
        await reset(dp, session, tmp, "s5")

        config = CHAIR_NORD.default_config()
        rows = [
            (
                CHAIR_NORD.id,
                CHAIR_NORD.name,
                CHAIR_NORD.config_line(config),
                config,
                1,
                42_900,
            )
        ] * 6
        for _ in range(10):
            await db.add_to_cart(USER_ID, CHAIR_NORD.id, config, 42_900)
            await db.create_order(USER_ID, "tester", ORDER_FORM, rows, 6 * 42_900)

        history = texts.orders_list(await db.user_orders(USER_ID, limit=10))
        check(
            f"10 заказов по 6 позиций -> {len(history)} знаков",
            len(history) <= texts.MESSAGE_LIMIT,
            len(history),
        )
        check("видно, что список обрезан", "Показаны последние" in history)

        # ------------------------------------------------------------------ 6
        print("\n[6] Имя клиента не превращается в формулу Excel")
        await reset(dp, session, tmp, "s6")

        evil = '=HYPERLINK("http://evil.tld";"Клиент")'
        await db.add_to_cart(USER_ID, CHAIR_NORD.id, config, 42_900)
        await db.create_order(
            USER_ID,
            "tester",
            {
                "name": evil,
                "phone": "+79001112233",
                "city": "-1+1",
                "address": "@somewhere",
                "comment": "+79001112233",
            },
            [
                (
                    CHAIR_NORD.id,
                    CHAIR_NORD.name,
                    CHAIR_NORD.config_line(config),
                    config,
                    1,
                    42_900,
                )
            ],
            42_900,
        )
        export_rows = await db.export_rows()

        sheet = load_workbook(io.BytesIO(export.to_xlsx(export_rows))).active
        formulas = [
            cell.value
            for row in sheet.iter_rows(min_row=2)
            for cell in row
            if cell.data_type == "f"
        ]
        check("в xlsx нет ни одной формулы", not formulas, formulas)

        csv_text = export.to_csv(export_rows).decode("utf-8-sig")
        dangerous = [
            field
            for line in csv_text.splitlines()[1:]
            for field in line.split(";")
            if field.startswith(("=", "+", "@"))
        ]
        check("в csv нет полей, начинающихся с = + @", not dangerous, dangerous)
        check("при этом данные не потеряны", "evil.tld" in csv_text)

        await bot.session.close()

    # ---------------------------------------------------------------------- 7
    print("\n[7] Время считается по Москве, а не по часам сервера")
    check("смещение +03:00", now().utcoffset().total_seconds() == 3 * 3600, now().utcoffset())
    check("метка зоны MSK", now().tzname() == "MSK", now().tzname())
    check("SHOP_TZ применён", now().tzinfo is SHOP_TZ)

    # ---------------------------------------------------------------------- 8
    print("\n[8] Тексты кнопок не разъехались между клавиатурой и фильтром")
    cancel_label = cancel_kb().keyboard[0][0].text
    skip_labels = [button.text for row in skip_kb().keyboard for button in row]
    check(
        "«Отменить» совпадает байт в байт",
        cancel_label == texts.BTN_CANCEL,
        f"{cancel_label!r} != {texts.BTN_CANCEL!r}",
    )
    check("«Пропустить» совпадает байт в байт", texts.BTN_SKIP in skip_labels, skip_labels)

    print("\n" + "=" * 46)
    if failed:
        print(f"ПРОВАЛЕНО {len(failed)} из {passed + len(failed)}:")
        for line in failed:
            print("  -", line)
        sys.exit(1)
    print(f"ВСЕ {passed} ПРОВЕРОК ПРОЙДЕНЫ")


if __name__ == "__main__":
    asyncio.run(main())
