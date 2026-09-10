"""Выгрузка заказов в Excel и CSV.

Файл собирается в памяти и сразу уходит в Telegram — на диске ничего не остаётся,
поэтому персональные данные клиентов не расползаются по серверу.
"""

from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .db import STATUSES
from .utils import format_dt

COLUMNS: tuple[tuple[str, str, int], ...] = (
    ("order_id", "№ заказа", 10),
    ("created_at", "Дата", 18),
    ("status", "Статус", 20),
    ("customer_name", "Клиент", 22),
    ("phone", "Телефон", 18),
    ("city", "Город", 18),
    ("address", "Адрес", 34),
    ("product_name", "Товар", 20),
    ("config_text", "Комплектация", 44),
    ("qty", "Кол-во", 8),
    ("unit_price", "Цена за шт., ₽", 15),
    ("order_total", "Сумма заказа, ₽", 16),
    ("comment", "Комментарий", 30),
    ("username", "Telegram", 18),
)

_HEADER_FILL = PatternFill("solid", fgColor="1F6F4A")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


# Excel считает формулой всё, что начинается с этих символов. Имя клиента вида
# =HYPERLINK("http://…"&A2;"Клиент") выполнится на компьютере администратора и
# может утащить телефоны и адреса других заказчиков из соседних ячеек.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _defuse(value: str) -> str:
    """Обезвреживает формулу апострофом: Excel покажет текст как есть."""
    return "'" + value if value.startswith(_FORMULA_PREFIXES) else value


def _pretty(key: str, value: object) -> object:
    if key == "status":
        result: object = STATUSES.get(str(value), value)
    elif key == "created_at":
        result = format_dt(str(value))
    elif key == "username" and value:
        # «@tester» Excel тоже считает началом формулы — обезвреживаем и его.
        result = f"@{value}"
    elif value is None:
        result = ""
    else:
        result = value
    # Единственный выход из функции: ни одна строка не минует обезвреживание.
    return _defuse(result) if isinstance(result, str) else result


def timestamped_name(extension: str) -> str:
    return f"nord-orders-{datetime.now():%Y%m%d-%H%M}.{extension}"


def to_xlsx(rows: list[dict[str, object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Заказы"

    for index, (_, title, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row_index, row in enumerate(rows, start=2):
        for col_index, (key, _, _) in enumerate(COLUMNS, start=1):
            sheet.cell(row=row_index, column=col_index, value=_pretty(key, row.get(key)))

    # Шапка закреплена и включён автофильтр — с таблицей сразу можно работать.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(COLUMNS))}{max(len(rows) + 1, 2)}"
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def to_csv(rows: list[dict[str, object]]) -> bytes:
    buffer = io.StringIO()
    # ';' и BOM — иначе русский Excel открывает CSV одной колонкой в кракозябрах.
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow([title for _, title, _ in COLUMNS])
    for row in rows:
        writer.writerow([_pretty(key, row.get(key)) for key, _, _ in COLUMNS])
    return buffer.getvalue().encode("utf-8-sig")
